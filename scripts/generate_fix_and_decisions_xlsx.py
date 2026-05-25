from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "docs" / "danh_sach_fix_va_huong_ky_thuat_bamboo_2026-04-19.xlsx"


FIX_ROWS = [
    {
        "stt": 1,
        "nhom": "Kiosk flow",
        "hang_muc": "Fix reset khi restart nhưng trên khay còn danh thiếp cũ",
        "van_de": "Card cũ còn trên khay sau restart làm phase REMOVING kẹt, hệ thống không reset về phiên mới.",
        "xu_ly": "Tách rõ hai nhánh REMOVING: hậu đăng ký bình thường và card cũ còn trên khay; chỉ nhánh hậu đăng ký mới chờ completion audio.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/mainpy.js",
        "moc": "14_4",
    },
    {
        "stt": 2,
        "nhom": "Face capture",
        "hang_muc": "Fix auto-capture không chạy sau khi nhận diện khách quay lại",
        "van_de": "Popup chào mừng quay lại xong thì guide không xanh, auto-capture không hoạt động.",
        "xu_ly": "Clear recognition hold và reset lại toàn bộ mốc FACE khi vào flow face thật sự.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/mainpy.js",
        "moc": "14_4",
    },
    {
        "stt": 3,
        "nhom": "Popup / Reset",
        "hang_muc": "Sửa điều kiện reset khi khách rút card trong lúc chờ cảm ơn / popup chức năng",
        "van_de": "Card bị rút sớm làm reset sai thời điểm hoặc kẹt popup.",
        "xu_ly": "Thêm cờ pending reset và chỉ reset sau khi popup chức năng hoặc popup đặt lịch thành công đã đóng.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/mainpy.js",
        "moc": "14_4",
    },
    {
        "stt": 4,
        "nhom": "QR print",
        "hang_muc": "Fix đã có registration_qr.png nhưng máy không in",
        "van_de": "App in qua file runtime trung gian lỗi đường dẫn, trong khi file gốc in tay được.",
        "xu_ly": "Đổi luồng in mặc định sang dùng trực tiếp registrations/<REG_ID>/registration_qr.png bằng đường dẫn tuyệt đối.",
        "ket_qua": "Đã hoàn thành",
        "files": "qr_printing/service.py, .env",
        "moc": "14_4",
    },
    {
        "stt": 5,
        "nhom": "OCR merge",
        "hang_muc": "Fix OCR đúng nhưng field fill sai hoặc bị mất dữ liệu",
        "van_de": "Kết quả OCR partial làm DB và state frontend overwrite mất field đúng.",
        "xu_ly": "Merge field theo nguyên tắc giữ giá trị cũ nếu giá trị mới rỗng; đồng bộ merge ở DB và frontend.",
        "ket_qua": "Đã hoàn thành",
        "files": "database/update_database.py, static/js/mainpy.js",
        "moc": "15_4",
    },
    {
        "stt": 6,
        "nhom": "Presence / Face",
        "hang_muc": "Fix vòng presence chết sau greeting khách quay lại",
        "van_de": "Nhánh greeting dùng return làm thoát loop Cam2, khiến guide và auto-capture mất dữ liệu live.",
        "xu_ly": "Bỏ return thoát loop, chỉ block greeting ở frame hiện tại và giữ presence stream tiếp tục chạy.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/mainpy.js",
        "moc": "15_4",
    },
    {
        "stt": 7,
        "nhom": "Q&A UI",
        "hang_muc": "Tách giao diện Q&A riêng khỏi kiosk chính",
        "van_de": "Q&A lẫn vào kiosk làm khó kiểm soát flow và khó mở rộng AI giai đoạn sau.",
        "xu_ly": "Tạo page /qa-voice-demo riêng với HTML/CSS/JS tách biệt, vẫn giữ liên kết từ popup chức năng.",
        "ket_qua": "Đã hoàn thành",
        "files": "templates/qa_voice_demo.html, static/js/qa_voice_demo.js, static/css/qa_voice_demo.css",
        "moc": "giai đoạn 4",
    },
    {
        "stt": 8,
        "nhom": "Q&A history",
        "hang_muc": "Đổi cách lưu Q&A theo phiên hội thoại",
        "van_de": "Lưu theo từng cặp hỏi-đáp làm dashboard khó đọc và không phản ánh đúng một phiên tương tác.",
        "xu_ly": "Thêm session_key, conversation_json, turn_count; append nhiều lượt vào một record cho đến khi làm mới.",
        "ket_qua": "Đã hoàn thành",
        "files": "database/update_database.py, routes/qa.py, static/js/qa_voice_demo.js, static/js/dashboard.js",
        "moc": "17_4",
    },
    {
        "stt": 9,
        "nhom": "Q&A voice",
        "hang_muc": "Fix trùng câu hỏi / trùng câu trả lời",
        "van_de": "SpeechRecognition submit lặp làm một câu hỏi sinh ra nhiều câu trả lời, trong đó có fallback giả.",
        "xu_ly": "Thêm dedupe transcript, dedupe submit và tách lỗi playback khỏi lỗi lấy answer text.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/qa_voice_demo.js",
        "moc": "RAG/QA",
    },
    {
        "stt": 10,
        "nhom": "Q&A latency",
        "hang_muc": "Rút ngắn thời gian phản hồi cho FAQ / câu hỏi factual",
        "van_de": "Các câu như Xin chào hoặc Smart Box dùng để làm gì bị chậm do đi qua RAG/LLM không cần thiết.",
        "xu_ly": "Thêm faq_extractive, extractive_fast, cache kết quả, cache retrieval, rút gọn context prompt.",
        "ket_qua": "Đã hoàn thành",
        "files": "services/qa_service.py",
        "moc": "RAG/QA",
    },
    {
        "stt": 11,
        "nhom": "RAG quality",
        "hang_muc": "Bổ sung eval retrieval và cải thiện route/source filtering",
        "van_de": "Retrieval dễ kéo sai nguồn hoặc trả lời bừa khi route / source specificity chưa đủ chặt.",
        "xu_ly": "Tạo bộ eval chuẩn, thêm source-priority filtering theo product và siết route rule/faq/rag/fallback.",
        "ket_qua": "Đã hoàn thành",
        "files": "rag/eval/qa_retrieval_eval_cases.json, scripts/run_qa_retrieval_eval.py, services/qa_service.py",
        "moc": "RAG/QA",
    },
    {
        "stt": 12,
        "nhom": "QR data",
        "hang_muc": "Đổi nội dung QR registration sang text nhiều dòng dễ đọc",
        "van_de": "QR chỉ chứa REG hoặc payload JSON kỹ thuật không phù hợp với nhu cầu quét xem thông tin trực tiếp.",
        "xu_ly": "Sinh QR text theo dạng ID / Họ tên / Công ty / Email / Điện thoại và thêm parser tương ứng ở kiosk.",
        "ket_qua": "Đã hoàn thành",
        "files": "services/registration.py, static/js/mainpy.js",
        "moc": "19_4",
    },
    {
        "stt": 13,
        "nhom": "QR print layout",
        "hang_muc": "Tách layout in và profile máy in ra file riêng",
        "van_de": "Căn chỉnh bằng .env khó kiểm soát, không phản ánh rõ phần bố cục ảnh và phần option driver.",
        "xu_ly": "Tạo layout.json cho bố cục tem và print_profile.json cho option lp/lpr; backend ưu tiên hai file này.",
        "ket_qua": "Đã hoàn thành",
        "files": "qr_printing/layout.json, qr_printing/print_profile.json, qr_printing/service.py",
        "moc": "19_4",
    },
    {
        "stt": 14,
        "nhom": "Dashboard",
        "hang_muc": "Tối ưu dashboard theo lazy load",
        "van_de": "Mở dashboard tải quá nhiều màn và refresh nền dày làm tốn request và render thừa.",
        "xu_ly": "Chỉ nạp stats, notifications và màn active; tăng chu kỳ polling; thêm guard chống request chồng.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/dashboard.js",
        "moc": "17_4",
    },
    {
        "stt": 15,
        "nhom": "Refactor",
        "hang_muc": "Giảm monolith của mainpy.js",
        "van_de": "mainpy.js quá lớn, dễ phát sinh race condition và regression khi sửa nhiều luồng khác nhau.",
        "xu_ly": "Tách timer, audio, welcome, async polling, presence thành helper modules và hard extraction khỏi file chính.",
        "ket_qua": "Đã hoàn thành",
        "files": "static/js/mainpy.js, static/js/kiosk_timers.js, static/js/kiosk_audio.js, static/js/kiosk_welcome.js, static/js/kiosk_async.js, static/js/kiosk_presence.js",
        "moc": "17_4",
    },
]


DECISION_ROWS = [
    {
        "stt": 1,
        "chu_de": "Kiến trúc Q&A",
        "quyet_dinh": "Tách Q&A khỏi kiosk chính, chạy trên /qa-voice-demo",
        "ly_do": "Giảm xung đột state với kiosk flow và tạo không gian riêng để phát triển AI hỏi đáp.",
        "pham_vi": "Frontend Q&A",
        "files": "templates/qa_voice_demo.html, static/js/qa_voice_demo.js, static/css/qa_voice_demo.css",
        "ghi_chu": "Popup chức năng chỉ điều hướng sang page QA, không nhúng QA trực tiếp vào index.",
    },
    {
        "stt": 2,
        "chu_de": "Nguồn dữ liệu RAG",
        "quyet_dinh": "Nguồn ingest chính thức của RAG là thư mục Rag_data",
        "ly_do": "Tách dữ liệu tri thức khỏi log sửa lỗi và dữ liệu runtime, thuận tiện rebuild và kiểm soát corpus.",
        "pham_vi": "Backend Q&A / RAG",
        "files": "Rag_data/, services/qa_service.py",
        "ghi_chu": "Log ngày và rag.md chỉ phục vụ theo dõi, không đưa vào corpus chính.",
    },
    {
        "stt": 3,
        "chu_de": "Router Q&A",
        "quyet_dinh": "Tách route trả lời thành rule / faq / rag / fallback",
        "ly_do": "Không để mọi câu đều đi qua LLM; giữ tính kiểm soát cho FAQ và workflow nội bộ.",
        "pham_vi": "Backend Q&A",
        "files": "services/qa_service.py",
        "ghi_chu": "Dynamic data được reserve riêng, không trộn vào RAG tĩnh.",
    },
    {
        "stt": 4,
        "chu_de": "Lưu lịch sử Q&A",
        "quyet_dinh": "Lưu theo phiên hội thoại, không lưu rời từng cặp hỏi-đáp",
        "ly_do": "Phù hợp với hành vi người dùng và giúp dashboard đọc lịch sử đúng ngữ cảnh hơn.",
        "pham_vi": "DB + Dashboard + QA UI",
        "files": "database/update_database.py, routes/qa.py, static/js/dashboard.js",
        "ghi_chu": "Chỉ mở phiên mới khi người dùng bấm Làm mới.",
    },
    {
        "stt": 5,
        "chu_de": "Voice command",
        "quyet_dinh": "Voice command chỉ áp dụng trong qa-voice-demo, vẫn giữ button tay fallback",
        "ly_do": "Giảm rủi ro tác động vào kiosk flow chính và tránh khóa UX nếu speech API lỗi.",
        "pham_vi": "Frontend Q&A",
        "files": "static/js/qa_voice_demo.js",
        "ghi_chu": "Button Bắt đầu / Kết thúc / Làm mới / Quay lại vẫn là đường fallback chính.",
    },
    {
        "stt": 6,
        "chu_de": "Cấu trúc kiosk",
        "quyet_dinh": "Giữ mainpy.js làm entry point nhưng tách các khối hạ tầng ra helper modules",
        "ly_do": "Giảm rủi ro refactor quá lớn trong khi vẫn bớt monolith và race condition.",
        "pham_vi": "Frontend kiosk",
        "files": "static/js/mainpy.js và các file kiosk_*.js",
        "ghi_chu": "Chọn cách refactor tăng dần thay vì rewrite toàn bộ.",
    },
    {
        "stt": 7,
        "chu_de": "QR registration",
        "quyet_dinh": "QR mới dùng text nhiều dòng dễ đọc, không dùng payload JSON kỹ thuật",
        "ly_do": "Phù hợp với nhu cầu quét trực tiếp bằng thiết bị ngoài kiosk và giảm tính kỹ thuật không cần thiết.",
        "pham_vi": "QR registration",
        "files": "services/registration.py, static/js/mainpy.js",
        "ghi_chu": "Vẫn giữ khả năng tương thích với tem cũ dạng REG và QR cũ đã in.",
    },
    {
        "stt": 8,
        "chu_de": "Cấu hình in QR",
        "quyet_dinh": "Tách bố cục ảnh và option driver thành hai file layout.json và print_profile.json",
        "ly_do": "Phân tách rõ phần ảnh cần in và phần option lp/lpr; giảm phụ thuộc .env.",
        "pham_vi": "QR printing",
        "files": "qr_printing/layout.json, qr_printing/print_profile.json, qr_printing/service.py",
        "ghi_chu": "layout.json cho size/offset ảnh; print_profile.json cho PageSize/orientation/scaling/options.",
    },
    {
        "stt": 9,
        "chu_de": "Profile máy in QR",
        "quyet_dinh": "Khóa profile in tự động theo Custom 78x52 mm, Landscape, Scale 254%",
        "ly_do": "Đây là bộ thông số thực tế đã kiểm tra thủ công cho kết quả in ổn định nhất.",
        "pham_vi": "QR printing",
        "files": "qr_printing/print_profile.json",
        "ghi_chu": "Luồng in tự động phải bám đúng profile này thay vì phụ thuộc hộp thoại Cmd/Win + P.",
    },
    {
        "stt": 10,
        "chu_de": "Nguồn ảnh in QR",
        "quyet_dinh": "In trực tiếp registration_qr.png đã dựng sẵn",
        "ly_do": "Giảm lỗi do render trung gian và giữ đúng bố cục ảnh đã chốt.",
        "pham_vi": "QR printing",
        "files": "services/registration.py, qr_printing/service.py",
        "ghi_chu": "Canvas runtime cũ chỉ giữ như legacy fallback.",
    },
    {
        "stt": 11,
        "chu_de": "Tối ưu dashboard",
        "quyet_dinh": "Dùng lazy load theo màn và refresh theo active screen",
        "ly_do": "Giảm tải nền, giảm request chồng và làm dashboard phản hồi nhẹ hơn.",
        "pham_vi": "Dashboard",
        "files": "static/js/dashboard.js",
        "ghi_chu": "Không bootstrap đồng loạt tất cả màn khi mở dashboard.",
    },
    {
        "stt": 12,
        "chu_de": "Định hướng RAG",
        "quyet_dinh": "Ưu tiên đúng dữ liệu và router trước khi tối ưu reranker/model",
        "ly_do": "Chất lượng retrieval và fallback quan trọng hơn việc thêm tối ưu model quá sớm.",
        "pham_vi": "Q&A / RAG",
        "files": "services/qa_service.py, rag/eval/qa_retrieval_eval_cases.json, rag/md/rag.md",
        "ghi_chu": "Thứ tự chốt là: corpus -> router -> eval -> tối ưu tiếp.",
    },
]


THIN = Side(style="thin", color="D0D7DE")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


def style_header(ws, row_idx: int, end_col: int) -> None:
    for col in range(1, end_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_table(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.alignment = WRAP_ALIGNMENT
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def autosize(ws, widths: dict[int, int]) -> None:
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_workbook() -> Workbook:
    wb = Workbook()
    ws_fix = wb.active
    ws_fix.title = "Danh_sach_fix"

    ws_fix["A1"] = "DANH SÁCH FIX ĐÃ THỰC HIỆN"
    ws_fix["A1"].font = Font(bold=True, size=14)
    ws_fix["A2"] = f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_fix["A3"] = "Phạm vi: Tổng hợp các lỗi/chỉnh sửa chính đã được xử lý trong dự án Bamboo Kiosk."
    for cell_ref in ("A1", "A2", "A3"):
        ws_fix[cell_ref].fill = TITLE_FILL if cell_ref == "A1" else PatternFill(fill_type=None)
    ws_fix.merge_cells("A1:H1")
    ws_fix.merge_cells("A2:H2")
    ws_fix.merge_cells("A3:H3")

    fix_headers = ["STT", "Nhóm", "Hạng mục / Fix", "Vấn đề trước khi sửa", "Hướng xử lý đã áp dụng", "Kết quả / Trạng thái", "File liên quan", "Mốc log"]
    header_row = 5
    for idx, header in enumerate(fix_headers, start=1):
        ws_fix.cell(row=header_row, column=idx, value=header)
    style_header(ws_fix, header_row, len(fix_headers))

    row = header_row + 1
    for item in FIX_ROWS:
        ws_fix.cell(row=row, column=1, value=item["stt"])
        ws_fix.cell(row=row, column=2, value=item["nhom"])
        ws_fix.cell(row=row, column=3, value=item["hang_muc"])
        ws_fix.cell(row=row, column=4, value=item["van_de"])
        ws_fix.cell(row=row, column=5, value=item["xu_ly"])
        ws_fix.cell(row=row, column=6, value=item["ket_qua"])
        ws_fix.cell(row=row, column=7, value=item["files"])
        ws_fix.cell(row=row, column=8, value=item["moc"])
        row += 1
    style_table(ws_fix, header_row + 1, row - 1, len(fix_headers))
    ws_fix.freeze_panes = "A6"
    ws_fix.auto_filter.ref = f"A5:H{row - 1}"
    autosize(ws_fix, {1: 6, 2: 18, 3: 38, 4: 42, 5: 48, 6: 18, 7: 42, 8: 10})

    ws_decision = wb.create_sheet("Huong_ky_thuat_chot")
    ws_decision["A1"] = "CÁC HƯỚNG KỸ THUẬT ĐÃ CHỐT"
    ws_decision["A1"].font = Font(bold=True, size=14)
    ws_decision["A2"] = f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_decision["A3"] = "Phạm vi: Tổng hợp các quyết định kỹ thuật đã được chốt để triển khai và bàn giao."
    ws_decision.merge_cells("A1:G1")
    ws_decision.merge_cells("A2:G2")
    ws_decision.merge_cells("A3:G3")
    ws_decision["A1"].fill = TITLE_FILL

    decision_headers = ["STT", "Chủ đề", "Quyết định kỹ thuật chốt", "Lý do chốt", "Phạm vi áp dụng", "File / Module chính", "Ghi chú"]
    header_row = 5
    for idx, header in enumerate(decision_headers, start=1):
        ws_decision.cell(row=header_row, column=idx, value=header)
    style_header(ws_decision, header_row, len(decision_headers))

    row = header_row + 1
    for item in DECISION_ROWS:
        ws_decision.cell(row=row, column=1, value=item["stt"])
        ws_decision.cell(row=row, column=2, value=item["chu_de"])
        ws_decision.cell(row=row, column=3, value=item["quyet_dinh"])
        ws_decision.cell(row=row, column=4, value=item["ly_do"])
        ws_decision.cell(row=row, column=5, value=item["pham_vi"])
        ws_decision.cell(row=row, column=6, value=item["files"])
        ws_decision.cell(row=row, column=7, value=item["ghi_chu"])
        row += 1
    style_table(ws_decision, header_row + 1, row - 1, len(decision_headers))
    ws_decision.freeze_panes = "A6"
    ws_decision.auto_filter.ref = f"A5:G{row - 1}"
    autosize(ws_decision, {1: 6, 2: 20, 3: 38, 4: 38, 5: 18, 6: 40, 7: 34})

    return wb


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
