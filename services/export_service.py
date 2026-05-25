"""
Export builders: XLSX and JSONL generation.
Extracted from application.py lines 1023-1224.
"""
import json
import logging
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, Border, Side

from app_config import logger, REG_DIR


def _jsonl_instruction_text() -> str:
    return "Trich xuat thong tin tu doan van ban OCR cua danh thiep sau. Tra ve dinh dang JSON."


def _detect_country(*values: str) -> str:
    haystack = " ".join(str(v or "") for v in values).lower()
    if "vietnam" in haystack or "viet nam" in haystack:
        return "Vietnam"
    if "japan" in haystack:
        return "Japan"
    if "singapore" in haystack:
        return "Singapore"
    if "thailand" in haystack:
        return "Thailand"
    if "china" in haystack:
        return "China"
    if "korea" in haystack:
        return "Korea"
    return ""


def build_jsonl_record(row: dict) -> dict:
    input_text = str(row.get("last_bcard_text") or "").strip()
    address = str(row.get("address") or "").strip()
    output_payload = {
        "company": str(row.get("company") or "").strip(),
        "full_name": str(row.get("full_name") or "").strip(),
        "job_title": str(row.get("title") or "").strip(),
        "phone": str(row.get("phone") or "").strip(),
        "email": str(row.get("email") or "").strip(),
        "address": address,
        "country": _detect_country(address, input_text),
    }
    return {
        "instruction": _jsonl_instruction_text(),
        "input": input_text,
        "output": output_payload,
    }


def build_jsonl_body(items: list[dict]) -> str:
    """Build JSONL export string from registration items."""
    lines = []
    for row in items:
        lines.append(json.dumps(build_jsonl_record(row), ensure_ascii=False))
    return "\n".join(lines)


def build_xlsx_bytes(items: list[dict]) -> bytes:
    """Build XLSX export bytes from registration items."""
    header = [
        "Registration ID",
        "Full Name",
        "Company",
        "Email",
        "Phone",
        "Job Title",
        "OCR Text",
        "Business Card Image",
        "Face Image",
        "QR Image",
        "Created At",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Database"
    ws.append(header)
    header_font = Font(bold=True)
    thin_side = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col in range(1, len(header) + 1):
        ws.cell(row=1, column=col).font = header_font
        ws.cell(row=1, column=col).border = cell_border

    ws.column_dimensions["H"].width = 22
    ws.column_dimensions["I"].width = 22
    ws.column_dimensions["J"].width = 22

    def _resolve_image_path(path_value, registration_id, image_stem):
        if not path_value:
            return None
        raw = str(path_value).strip()
        p = None

        if raw.startswith("/registrations/"):
            p = (Path.cwd() / raw.lstrip("/")).resolve()
        else:
            if "/registrations/" in raw.replace("\\", "/"):
                rel = raw.replace("\\", "/").split("/registrations/", 1)[1]
                p = (REG_DIR / rel).resolve()
            else:
                p = Path(raw)
            if not p.exists():
                base_dir = REG_DIR / str(registration_id or "")
                if p.name:
                    by_name = base_dir / p.name
                    if by_name.exists():
                        p = by_name
                if not p.exists():
                    for ext in ("jpeg", "jpg", "png", "webp"):
                        by_stem = base_dir / f"{image_stem}.{ext}"
                        if by_stem.exists():
                            p = by_stem
                            break

        if not p or not p.exists():
            return None
        return p

    def add_image(path_value, cell_ref, registration_id, image_stem):
        p = _resolve_image_path(path_value, registration_id, image_stem)
        if p is None:
            return
        try:
            img = OpenpyxlImage(str(p))
            max_w, max_h = 140, 100
            scale = min(max_w / img.width, max_h / img.height, 1.0)
            img.width = int(img.width * scale)
            img.height = int(img.height * scale)
            ws.add_image(img, cell_ref)
        except Exception:
            logger.exception("Cannot add image to excel: %s", p)

    for idx, row in enumerate(items, start=2):
        ws.cell(row=idx, column=1, value=str(row.get("registration_id", "")))
        ws.cell(row=idx, column=2, value=str(row.get("full_name", "")))
        ws.cell(row=idx, column=3, value=str(row.get("company", "")))
        ws.cell(row=idx, column=4, value=str(row.get("email", "")))
        ws.cell(row=idx, column=5, value=str(row.get("phone", "")))
        ws.cell(row=idx, column=6, value=str(row.get("title", "")))
        ws.cell(row=idx, column=7, value=str(row.get("last_bcard_text", "")))
        ws.cell(row=idx, column=11, value=str(row.get("created_at", "")))
        ws.row_dimensions[idx].height = 80

        reg_id = row.get("registration_id")
        add_image(row.get("bcard_link"), f"H{idx}", reg_id, "bcard")
        add_image(row.get("face_link"), f"I{idx}", reg_id, "face")
        add_image(row.get("qr_link"), f"J{idx}", reg_id, "registration_qr")

    last_row = max(1, len(items) + 1)
    for r in range(2, last_row + 1):
        for c in range(1, len(header) + 1):
            ws.cell(row=r, column=c).border = cell_border

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def build_cccd_json_body(items: list[dict]) -> str:
    return json.dumps(items, ensure_ascii=False, indent=2)


def build_cccd_xlsx_bytes(items: list[dict]) -> bytes:
    header = [
        "Registration ID",
        "ID Number",
        "Old ID",
        "Full Name",
        "DOB",
        "Gender",
        "Address",
        "Issued",
        "Expiry",
        "QR Raw",
        "OCR Text",
        "CCCD Front",
        "CCCD Back",
        "Face Image",
        "Created At",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "CCCD"
    ws.append(header)
    header_font = Font(bold=True)
    thin_side = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col in range(1, len(header) + 1):
        ws.cell(row=1, column=col).font = header_font
        ws.cell(row=1, column=col).border = cell_border

    ws.column_dimensions["L"].width = 22
    ws.column_dimensions["M"].width = 22
    ws.column_dimensions["N"].width = 22

    def _resolve_image_path(path_value, registration_id, image_stem):
        if not path_value:
            return None
        raw = str(path_value).strip()
        if raw.startswith("/registrations/"):
            p = (Path.cwd() / raw.lstrip("/")).resolve()
            return p if p.exists() else None
        p = Path(raw)
        if p.exists():
            return p
        base_dir = REG_DIR / str(registration_id or "")
        for ext in ("jpeg", "jpg", "png", "webp"):
            candidate = base_dir / f"{image_stem}.{ext}"
            if candidate.exists():
                return candidate
        return None

    def add_image(path_value, cell_ref, registration_id, image_stem):
        p = _resolve_image_path(path_value, registration_id, image_stem)
        if p is None:
            return
        try:
            img = OpenpyxlImage(str(p))
            max_w, max_h = 140, 100
            scale = min(max_w / img.width, max_h / img.height, 1.0)
            img.width = int(img.width * scale)
            img.height = int(img.height * scale)
            ws.add_image(img, cell_ref)
        except Exception:
            logger.exception("Cannot add image to excel: %s", p)

    for idx, row in enumerate(items, start=2):
        values = [
            row.get("registration_id", ""),
            row.get("id_number", ""),
            row.get("old_id", ""),
            row.get("full_name", ""),
            row.get("dob", ""),
            row.get("gender", ""),
            row.get("address", ""),
            row.get("issued", ""),
            row.get("expiry", ""),
            row.get("cccd_qr_raw", ""),
            row.get("cccd_ocr_text", ""),
            "",
            "",
            "",
            row.get("created_at", ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=idx, column=col, value=str(value))
        ws.row_dimensions[idx].height = 80
        reg_id = row.get("registration_id")
        add_image(row.get("cccd_front_link"), f"L{idx}", reg_id, "cccd_front")
        add_image(row.get("cccd_back_link"), f"M{idx}", reg_id, "cccd_back")
        add_image(row.get("face_link"), f"N{idx}", reg_id, "face")

    last_row = max(1, len(items) + 1)
    for r in range(2, last_row + 1):
        for c in range(1, len(header) + 1):
            ws.cell(row=r, column=c).border = cell_border

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
